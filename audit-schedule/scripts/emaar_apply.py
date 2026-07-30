"""Build and write the August 2026 EMAAR AUDT schedule.

Rule chosen by the user: the PPM date wins.  Where a zone has any Aug PM date other
than the one its Joint Audit uses, the Emaar audit goes on one of those (the furthest
from the Joint Audit).  Where the PPM offers no alternative -- 16 zones have a single
PM date all month, 33 more are absent from the PPM entirely -- the audit falls on a
free day at least a week from the Joint Audit.
"""
import json, copy, re, openpyxl
from collections import defaultdict, Counter
from openpyxl.utils import get_column_letter as gcl

BOOK = 'Audit_Schedule_JuLY_2026_TRANSGUARD_DM.xlsx'
FIRST, LAST, TOT = 5, 88, 89
DAY0, CAP, GAP = 2, 4, 7
SLOT = [0, 1, 4, 9]          # convex penalty on the 1st..4th audit of a day

# ---- Aug PM dates per zone, rebuilt from the PPM -----------------------------
p = openpyxl.load_workbook('aug_ppm.xlsx')['List of PMs']
ALIAS = {'Z01': 'Z1', 'Z01A': 'Z01a', 'E33': 'Z33'}
pm = defaultdict(set)
for r in range(2, p.max_row + 1):
    if not p.cell(r, 1).value: continue
    d = (p.cell(r, 2).value or '').replace('\xa0', ' ').strip()
    tok = re.search(r'((?:ZB|FV|CT|Z|E)\s*0*\d+[a-zA-Z]?)\s*$', d).group(1).replace(' ', '').upper()
    m, day, y = (int(x) for x in p.cell(r, 3).value.split('/'))
    if (m, y) == (8, 26): pm[ALIAS.get(tok, tok).upper()].add(day)

wb = openpyxl.load_workbook(BOOK)
ja_ws, ws = wb['JOINT AUDIT '], wb['EMAAR AUDT']

zones = [ws.cell(r, 1).value for r in range(FIRST, LAST + 1)]
zrow = {z: FIRST + i for i, z in enumerate(zones)}
JA = {z: next(c - DAY0 for c in range(3, 34) if ja_ws.cell(zrow[z], c).value == 'AU') for z in zones}
assert [ja_ws.cell(r, 1).value for r in range(FIRST, LAST + 1)] == zones, 'sheet rows are misaligned'

# ---- candidate days ----------------------------------------------------------
edges, pm_anchored = {}, set()
for z in zones:
    alts = sorted(d for d in pm.get(z.upper(), ()) if d != JA[z])
    if alts:                                   # PPM offers a different date -> use it
        pm_anchored.add(z)
        # Push every gap towards a full week first (squared shortfall dominates), then
        # maximise the remaining gap.  Both outrank the load-balancing slot costs, so
        # spreading the month can never buy itself a day of separation.
        edges[z] = {d: 1000 * max(0, GAP - abs(d - JA[z])) ** 2 + 20 * (31 - abs(d - JA[z]))
                    for d in alts}
    else:                                      # no alternative -> free day, >= a week away
        edges[z] = {d: 0 for d in range(1, 32) if abs(d - JA[z]) >= GAP}

# ---- min-cost max-flow -------------------------------------------------------
class MCMF:
    def __init__(s, n): s.n, s.g = n, [[] for _ in range(n)]
    def add(s, u, v, cap, cost):
        s.g[u].append([v, cap, cost, len(s.g[v])]); s.g[v].append([u, 0, -cost, len(s.g[u]) - 1])
    def run(s, src, snk):
        INF = float('inf'); flow = cost = 0
        while True:
            dist = [INF]*s.n; dist[src] = 0; inq = [False]*s.n
            pv = [-1]*s.n; pe = [-1]*s.n; q = [src]; inq[src] = True
            while q:
                u = q.pop(0); inq[u] = False
                for i, e in enumerate(s.g[u]):
                    v, cap, c, _ = e
                    if cap > 0 and dist[u] + c < dist[v]:
                        dist[v] = dist[u] + c; pv[v], pe[v] = u, i
                        if not inq[v]: q.append(v); inq[v] = True
            if dist[snk] == INF: return flow, cost
            d = INF; v = snk
            while v != src: d = min(d, s.g[pv[v]][pe[v]][1]); v = pv[v]
            v = snk
            while v != src:
                e = s.g[pv[v]][pe[v]]; e[1] -= d; s.g[v][e[3]][1] += d; v = pv[v]
            flow += d; cost += d*dist[snk]

n = len(zones); S = 0; T = n + 31 + 1
zi = {z: i + 1 for i, z in enumerate(zones)}; di = {d: n + d for d in range(1, 32)}
mc = MCMF(T + 1)
for z in zones:
    mc.add(S, zi[z], 1, 0)
    for d, c in edges[z].items(): mc.add(zi[z], di[d], 1, c)
for d in range(1, 32):
    for sc in SLOT: mc.add(di[d], T, 1, sc)
flow, _ = mc.run(S, T)
assert flow == n, f'infeasible: {flow}/{n}'

EA = {}
for z in zones:
    for v, cap, c, _ in mc.g[zi[z]]:
        if cap == 0 and v in di.values():
            EA[z] = next(d for d, i in di.items() if i == v); break

# ---- write -------------------------------------------------------------------
model = ja_ws['F5']
font, align, fill = copy.copy(model.font), copy.copy(model.alignment), copy.copy(model.fill)
align.vertical = 'center'
for r in range(FIRST, LAST + 1):
    for c in range(3, 34):
        if ws.cell(r, c).value is not None: ws.cell(r, c).value = None
for z in zones:
    cell = ws.cell(zrow[z], EA[z] + DAY0)
    cell.value = 'AU'
    cell.font, cell.alignment, cell.fill = copy.copy(font), copy.copy(align), copy.copy(fill)
for d in range(1, 32):                      # AF89 was a hardcoded 0, AG89 pointed at empty AG93
    L = gcl(d + DAY0)
    ws.cell(TOT, d + DAY0).value = f'=COUNTA({L}{FIRST}:{L}{LAST})'

wb.calculation.fullCalcOnLoad = True
wb.save(BOOK)

load = Counter(EA.values()); gaps = {z: abs(EA[z] - JA[z]) for z in zones}
print(f'EMAAR audits on a PM date : {len(pm_anchored)} of 84')
print(f'free-day placements       : {84 - len(pm_anchored)} (16 single-PM-date zones + 33 not in the PPM)')
print(f'same day as Joint Audit   : {sum(1 for z in zones if EA[z] == JA[z])}')
print(f'gap: min {min(gaps.values())}, median {sorted(gaps.values())[42]}, max {max(gaps.values())}')
print(f'busiest day: {max(load.values())} audits (cap {CAP})')
json.dump({'EA': EA, 'JA': JA, 'pm_anchored': sorted(pm_anchored)}, open('emaar_final.json', 'w'))
