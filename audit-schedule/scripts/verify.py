"""Independent end-to-end check of the delivered workbook."""
import openpyxl, re
from collections import defaultdict, Counter
from openpyxl.utils import get_column_letter as gcl, range_boundaries

SRC, OUT = 'audit_july.xlsx', 'Audit_Schedule_JuLY_2026_TRANSGUARD_DM.xlsx'
fails = []
def check(ok, msg):
    print(('  PASS  ' if ok else '  FAIL  ') + msg)
    if not ok: fails.append(msg)

# ---- truth table rebuilt from the PPM, independent of the solver -------------
p = openpyxl.load_workbook('aug_ppm.xlsx')['List of PMs']
ALIAS = {'Z01': 'Z1', 'Z01A': 'Z01A', 'E33': 'Z33'}
truth, ppm_rows = defaultdict(set), 0
for r in range(2, p.max_row + 1):
    if not p.cell(r, 1).value: continue
    ppm_rows += 1
    d = (p.cell(r, 2).value or '').replace('\xa0', ' ').strip()
    tok = re.search(r'((?:ZB|FV|CT|Z|E)\s*0*\d+[a-zA-Z]?)\s*$', d).group(1).replace(' ', '').upper()
    m, day, y = (int(x) for x in p.cell(r, 3).value.split('/'))
    if (m, y) == (8, 26): truth[ALIAS.get(tok, tok)].add(day)

src, out = openpyxl.load_workbook(SRC), openpyxl.load_workbook(OUT)
so, oo = src['JOINT AUDIT '], out['JOINT AUDIT ']

print('== JOINT AUDIT schedule ==')
sched, load = {}, Counter()
one_each, in_grid = True, True
for r in range(5, 89):
    z = oo.cell(r, 1).value
    marks = [c - 2 for c in range(3, 34) if oo.cell(r, c).value == 'AU']
    if len(marks) != 1: one_each = False
    if any(oo.cell(r, c).value == 'AU' for c in range(34, 40)): in_grid = False
    if marks: sched[z] = marks[0]; load[marks[0]] += 1
check(one_each, 'every one of the 84 zones is marked exactly once')
check(in_grid, 'no audit mark falls outside the day columns C..AG')
check(sum(load.values()) == 84, f'total audits = {sum(load.values())} (expected 84)')
check(all(1 <= d <= 31 for d in sched.values()), 'every audit lands on a day within August (1-31)')
check(max(load.values()) <= 4, f'no day exceeds 4 audits (busiest day has {max(load.values())})')

bad = [f'{z}->{d} not in {sorted(truth[z.upper()])}' for z, d in sched.items()
       if z.upper() in truth and d not in truth[z.upper()]]
covered = [z for z in sched if z.upper() in truth]
check(not bad, f'all {len(covered)} PPM-driven zones sit on one of their own Aug PM dates' + (f' :: {bad}' if bad else ''))

# zones with no Aug PPM entry must keep their July day
sj = {so.cell(r, 1).value: next(c - 2 for c in range(3, 34) if so.cell(r, c).value == 'AU') for r in range(5, 89)}
free = [z for z in sched if z.upper() not in truth]
check(all(sched[z] == sj[z] for z in free), f'all {len(free)} zones absent from the PPM kept their July day')
check(len(covered) + len(free) == 84, f'zone accounting: {len(covered)} PPM + {len(free)} non-PPM = 84')

print('\n== preserved content ==')
check([oo.cell(r,1).value for r in range(5,89)] == [so.cell(r,1).value for r in range(5,89)], 'zone column unchanged')
check([oo.cell(r,2).value for r in range(5,89)] == [so.cell(r,2).value for r in range(5,89)], 'auditor names unchanged')
check([oo.cell(4,c).value for c in range(1,40)] == [so.cell(4,c).value for c in range(1,40)], 'header row unchanged')
check(oo['A1'].value == so['A1'].value and oo['C3'].value == so['C3'].value, 'titles unchanged')
check(str(oo.merged_cells.ranges) == str(so.merged_cells.ranges), 'merged ranges unchanged')
check(oo.freeze_panes == so.freeze_panes and oo.auto_filter.ref == so.auto_filter.ref, 'freeze panes / autofilter unchanged')
for sh in ['INSTRUCTION', 'EMAAR AUDT', 'ZONE']:
    a = [[c.value for c in row] for row in src[sh].iter_rows()]
    b = [[c.value for c in row] for row in out[sh].iter_rows()]
    check(a == b, f"sheet '{sh}' untouched")
check(out.sheetnames == src.sheetnames, 'sheet names/order unchanged')

print('\n== formulas ==')
def counta(ws, rng):
    c1, r1, c2, r2 = range_boundaries(rng)
    return sum(1 for rr in range(r1, r2+1) for cc in range(c1, c2+1)
               if ws.cell(rr, cc).value not in (None, ''))
def summ(ws, rng, rowtot):
    c1, r1, c2, r2 = range_boundaries(rng)
    return sum(rowtot.get((rr, cc), 0) for rr in range(r1, r2+1) for cc in range(c1, c2+1))

# COUNTA/SUM over a range, or a bare cell reference (EMAAR!AG89 '=AG93' is pre-existing)
allowed_fn = re.compile(r'^=((COUNTA|SUM)\([A-Z]+\d+:[A-Z]+\d+\)|[A-Z]+\d+)$')
formulas, rowtot = [], {}
for ws in out.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                formulas.append((ws.title, c.coordinate, c.value))
check(all(allowed_fn.match(f) for _, _, f in formulas),
      f'all {len(formulas)} formulas are plain COUNTA/SUM over a single range (no volatile or post-2007 functions)')

# row totals then column totals
ok_rows = True
for r in range(5, 89):
    f = oo.cell(r, 35).value
    v = counta(oo, f[len('=COUNTA('):-1])
    rowtot[(r, 35)] = v
    if v != 1: ok_rows = False
check(ok_rows, 'AI column (per-zone total) evaluates to 1 for all 84 zones')

expect = {d: load.get(d, 0) for d in range(1, 32)}
got = {}
ok_cols = True
for d in range(1, 32):
    cell = oo.cell(89, d + 2)
    if not (isinstance(cell.value, str) and cell.value.startswith('=COUNTA(')):
        ok_cols = False; continue
    got[d] = counta(oo, cell.value[len('=COUNTA('):-1])
check(ok_cols, 'all 31 day columns have a COUNTA total (the two hardcoded 4s and the missing one are now formulas)')
check(got == expect, 'every day total evaluates to the actual number of audits on that day')
check(summ(oo, oo.cell(89,35).value[len('=SUM('):-1], rowtot) == 84, 'grand total AI89 evaluates to 84')

print('\n== summary ==')
print(f'PPM rows read: {ppm_rows} | zones with Aug PM dates: {len(truth)} | audits placed: {sum(load.values())}')
print('per-day audit count:', {d: load.get(d, 0) for d in range(1, 32)})
print('\nRESULT:', 'ALL CHECKS PASSED' if not fails else f'{len(fails)} FAILURE(S): {fails}')
