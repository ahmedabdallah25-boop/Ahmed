"""Build the August 2026 JOINT AUDIT schedule from the Aug 2026 PPM."""
import openpyxl, re, json
from collections import defaultdict, Counter

PPM = 'aug_ppm.xlsx'
AUD = 'audit_july.xlsx'
CAP = 4                      # max audits per day
SLOT_COST = [0, 1, 4, 9]     # convex penalty on the 1st..4th audit of a day
MOVE_BASE = 30               # cost of moving a non-PPM zone off its July day

# ---------------------------------------------------------------- PPM parse
ws = openpyxl.load_workbook(PPM)['List of PMs']
ppm_days = defaultdict(set)
for r in range(2, ws.max_row + 1):
    if not ws.cell(r, 1).value:
        continue
    desc = (ws.cell(r, 2).value or '').replace('\xa0', ' ').strip()
    tok = re.search(r'((?:ZB|FV|CT|Z|E)\s*0*\d+[a-zA-Z]?)\s*$', desc).group(1).replace(' ', '').upper()
    m, d, y = (int(x) for x in ws.cell(r, 3).value.split('/'))
    if (m, y) == (8, 26):
        ppm_days[tok].add(d)

ALIAS = {'Z01': 'Z1', 'Z01A': 'Z01a', 'E33': 'Z33'}   # PPM tag -> audit-sheet zone label

# ---------------------------------------------------------- audit sheet read
wa = openpyxl.load_workbook(AUD)['JOINT AUDIT ']
zone_row, july_day = {}, {}
for r in range(5, 89):
    z = wa.cell(r, 1).value
    zone_row[z] = r
    for c in range(3, 34):
        if wa.cell(r, c).value == 'AU':
            july_day[z] = c - 2
by_upper = {z.upper(): z for z in zone_row}

allowed = {}
for tag, days in ppm_days.items():
    z = by_upper[ALIAS.get(tag, tag).upper()]
    allowed.setdefault(z, set()).update(days)

zones = list(zone_row)
pm_zones = [z for z in zones if z in allowed]
free_zones = [z for z in zones if z not in allowed]

# ------------------------------------------------- min-cost max-flow (SSP)
class MCMF:
    def __init__(self, n):
        self.n, self.g = n, [[] for _ in range(n)]
    def add(self, u, v, cap, cost):
        self.g[u].append([v, cap, cost, len(self.g[v])])
        self.g[v].append([u, 0, -cost, len(self.g[u]) - 1])
    def run(self, s, t):
        INF = float('inf')
        flow = cost = 0
        while True:
            dist = [INF] * self.n; dist[s] = 0
            inq = [False] * self.n; prevv = [-1] * self.n; preve = [-1] * self.n
            queue = [s]; inq[s] = True
            while queue:                      # SPFA
                u = queue.pop(0); inq[u] = False
                for i, e in enumerate(self.g[u]):
                    v, cap, c, _ = e
                    if cap > 0 and dist[u] + c < dist[v]:
                        dist[v] = dist[u] + c
                        prevv[v], preve[v] = u, i
                        if not inq[v]:
                            queue.append(v); inq[v] = True
            if dist[t] == INF:
                return flow, cost
            d = INF; v = t
            while v != s:
                d = min(d, self.g[prevv[v]][preve[v]][1]); v = prevv[v]
            v = t
            while v != s:
                e = self.g[prevv[v]][preve[v]]
                e[1] -= d; self.g[v][e[3]][1] += d; v = prevv[v]
            flow += d; cost += d * dist[t]

S, T = 0, len(zones) + 31 + 1
zidx = {z: i + 1 for i, z in enumerate(zones)}
didx = {d: len(zones) + d for d in range(1, 32)}
mc = MCMF(T + 1)
for z in zones:
    mc.add(S, zidx[z], 1, 0)
    if z in allowed:
        for d in sorted(allowed[z]):
            mc.add(zidx[z], didx[d], 1, 0)
    else:
        jd = july_day[z]
        for d in range(1, 32):
            mc.add(zidx[z], didx[d], 1, 0 if d == jd else MOVE_BASE + abs(d - jd))
for d in range(1, 32):
    for sc in SLOT_COST:
        mc.add(didx[d], T, 1, sc)

flow, cost = mc.run(S, T)
assert flow == len(zones), f'infeasible: only {flow}/{len(zones)} zones placed'

assign = {}
for z in zones:
    for v, cap, c, _ in mc.g[zidx[z]]:
        if cap == 0 and v in didx.values():
            assign[z] = next(d for d, i in didx.items() if i == v)
            break

# --------------------------------------------------------------- verify
assert set(assign) == set(zones)
for z in pm_zones:
    assert assign[z] in allowed[z], (z, assign[z], allowed[z])
load = Counter(assign.values())
assert max(load.values()) <= CAP, load
moved = {z: (july_day[z], assign[z]) for z in free_zones if assign[z] != july_day[z]}

print(f'placed {len(assign)} zones, cost={cost}')
print(f'PPM-driven zones: {len(pm_zones)}   non-PPM zones: {len(free_zones)}')
print(f'max per day: {max(load.values())}   days used: {len(load)}')
print('per-day load:', dict(sorted(load.items())))
print('non-PPM zones moved off their July day:', moved or 'none')
json.dump({'assign': assign, 'allowed': {k: sorted(v) for k, v in allowed.items()},
           'july': july_day, 'row': zone_row}, open('assignment.json', 'w'), indent=1)
